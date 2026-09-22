from concurrent.futures import ThreadPoolExecutor, wait
from copy import copy
import json
import os
import re
import threading
import time
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from PIL import Image, ImageTk
import openpyxl
from openpyxl.utils import get_column_letter
import pandas as pd
from seniatOCR import SeniatOCR
from tkinterdnd2 import DND_FILES, TkinterDnD


class ExcelUploaderApp(TkinterDnD.Tk):

  def __init__(self):
    super().__init__()
    self.title("Procesador de RIFs SENIAT")
    self.geometry(self.cargar_configuracion())
    self.config(bg="#f8f9fa")
    self.resizable(True, True)
    self.minsize(800, 600)
    self.protocol("WM_DELETE_WINDOW", self.cerrar_app)

    self.df_filtrado_global = None
    self.ruta_archivo_actual = None
    self.route_vars = {}
    self.cancelar_proceso = False
    self.ruta_guardado = None

    self.chk_activo_var = tk.BooleanVar(value=True)
    self.chk_inactivo_var = tk.BooleanVar(value=True)

    self.chk_activo_var.trace_add(
        "write", lambda *args: self.actualizar_estadisticas()
    )
    self.chk_inactivo_var.trace_add(
        "write", lambda *args: self.actualizar_estadisticas()
    )

    # --- FRAME CABECERA (LOGO OPCIONAL Y ADAPTABLE) ---
    self.frame_logo = tk.Frame(self, bg="#f8f9fa")
    self.frame_logo.pack(side="top", anchor="nw", fill="x", padx=15, pady=10)

    try:
      script_dir = os.path.dirname(os.path.abspath(__file__))
      parent_dir = os.path.dirname(script_dir)

      posibles_nombres = [
          "LOGO_ROCO_3.jpg",
          "LOGO_ROCO_3.png",
          "LOGO_ROCO_3.jpeg",
          "logo_roco_3.jpg",
          "logo_roco_3.png",
      ]
      logo_path = None

      for directorio in [script_dir, parent_dir]:
        for nombre in posibles_nombres:
          ruta_prueba = os.path.join(directorio, nombre)
          if os.path.exists(ruta_prueba):
            logo_path = ruta_prueba
            break
        if logo_path:
          break

      if logo_path:
        img_pil = Image.open(logo_path).convert("RGBA")
        basewidth = 90
        wpercent = basewidth / float(img_pil.size[0])
        hsize = int(float(img_pil.size[1]) * wpercent)
        img_pil = img_pil.resize(
            (basewidth, hsize), Image.Resampling.LANCZOS
        )

        data = img_pil.getdata()
        new_data = []
        for item in data:
          if item[0] > 240 and item[1] > 240 and item[2] > 240:
            new_data.append((248, 249, 250, 255))
          else:
            new_data.append(item)
        img_pil.putdata(new_data)

        self.logo_img = ImageTk.PhotoImage(img_pil)

        lbl_logo = tk.Label(
            self.frame_logo, image=self.logo_img, bg="#f8f9fa", cursor="hand2"
        )
        lbl_logo.pack(side="left")
    except Exception:
      pass

    # --- FRAME DE CARGA INICIAL ---
    self.frame_upload = tk.Frame(self, bg="#f8f9fa")
    self.frame_upload.pack(expand=True, fill="both", padx=30, pady=10)

    self.btn_browse = tk.Button(
        self.frame_upload,
        text="Buscar en archivos",
        command=self.browse_file,
        font=("Arial", 10, "bold"),
        bg="#007bff",
        fg="white",
        padx=15,
        pady=8,
        relief="flat",
        cursor="hand2",
    )
    self.btn_browse.pack(side="bottom", pady=(0, 10))

    self.label = tk.Label(
        self.frame_upload,
        text="📁 Arrastra tu archivo Excel aquí\no usa el botón para buscarlo",
        bg="#f8f9fa",
        fg="#495057",
        font=("Arial", 11),
        justify="center",
    )
    self.label.pack(expand=True, fill="both")

    # --- FRAME INFO ARCHIVO CARGADO ---
    self.frame_file_info = tk.Frame(self, bg="#e9ecef", bd=1, relief="solid")

    self.lbl_file_name = tk.Label(
        self.frame_file_info,
        text="",
        bg="#e9ecef",
        fg="#0d6efd",
        font=("Arial", 9, "bold", "underline"),
        wraplength=320,
        justify="left",
        cursor="hand2",
    )
    self.lbl_file_name.pack(side="left", padx=8, pady=6)
    self.lbl_file_name.bind("<Button-1>", self.abrir_archivo_excel)

    self.btn_change_file = tk.Button(
        self.frame_file_info,
        text="🔄 Cambiar archivo",
        command=self.browse_file,
        font=("Arial", 8, "bold"),
        bg="#dc3545",
        fg="white",
        padx=6,
        pady=3,
        relief="flat",
        cursor="hand2",
    )
    self.btn_change_file.pack(side="right", padx=8, pady=6)

    # --- FRAME CONTENEDOR DE CHECKBOXES Y RUTAS ---
    self.frame_checks = tk.Frame(self, bg="#f8f9fa")
    self.frame_checks.pack(pady=2, fill="both", expand=True)

    # --- BOTÓN INICIAR CONSULTA ---
    self.btn_download = tk.Button(
        self,
        text="Iniciar Consulta",
        command=self.iniciar_hilo_consultas,
        font=("Arial", 10, "bold"),
        bg="#28a745",
        fg="white",
        padx=12,
        pady=5,
        relief="flat",
        cursor="hand2",
    )

    self.drop_target_register(DND_FILES)
    self.dnd_bind("<<Drop>>", self.handle_drop)

  def cargar_configuracion(self):
    config_path = os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "config.json"
    )
    if os.path.exists(config_path):
      try:
        with open(config_path, "r") as f:
          data = json.load(f)
          ancho = max(data.get("ancho", 500), 480)
          alto = max(data.get("alto", 410), 410)
          return f"{ancho}x{alto}"
      except Exception:
        pass
    return "500x410"

  def cerrar_app(self):
    try:
      geom = self.geometry()
      match = re.match(r"^(\d+)x(\d+)", geom)
      if match:
        ancho, alto = match.groups()
        config_path = os.path.join(
            os.path.dirname(os.path.abspath(__file__)), "config.json"
        )
        with open(config_path, "w") as f:
          json.dump({"ancho": int(ancho), "alto": int(alto)}, f)
    except Exception:
      pass
    self.destroy()

  def abrir_archivo_excel(self, event=None):
    if self.ruta_archivo_actual and os.path.exists(self.ruta_archivo_actual):
      try:
        os.startfile(self.ruta_archivo_actual)
      except Exception as e:
        messagebox.showerror("Error", f"No se pudo abrir el archivo:\n{e}")

  def browse_file(self):
    file_path = filedialog.askopenfilename(
        title="Seleccionar archivo Excel",
        filetypes=[
            ("Archivos de Excel", "*.xlsx"),
            ("Todos los archivos", "*.*"),
        ],
    )
    if file_path:
      self.process_excel(file_path)

  def handle_drop(self, event):
    file_path = event.data.strip("{}")
    if file_path.lower().endswith(".xlsx"):
      self.process_excel(file_path)
    else:
      messagebox.showerror(
          "Error", "Por favor, selecciona un archivo Excel válido (.xlsx)"
      )

  def process_excel(self, file_path):
    try:
      df_preview = pd.read_excel(file_path, header=None, nrows=30)
      header_idx = 8
      for idx, row in df_preview.iterrows():
        if any(
            "Cliente (Identificación)" in str(val) for val in row.values
        ):
          header_idx = idx
          break

      self.excel_header_row = header_idx

      df = pd.read_excel(file_path, header=header_idx)
      df = df.loc[:, ~df.columns.duplicated()]
      df.columns = [str(c).strip() for c in df.columns]

      col_ruta, col_rif, col_nombre, col_estatus = None, None, None, None

      for col in df.columns:
        col_upper = str(col).upper().strip()
        if not col_ruta and col_upper in ["RUTA (NOMBRE)", "RUTA"]:
          col_ruta = col
        elif not col_rif and col_upper in [
            "CLIENTE (IDENTIFICACIÓN)",
            "CLIENTE (IDENTIFICACION)",
            "RIF",
            "IDENTIFICACIÓN",
            "IDENTIFICACION",
        ]:
          col_rif = col
        elif not col_nombre and col_upper in [
            "NOMBRE CLIENTE",
            "CLIENTE (NOMBRE)",
        ]:
          col_nombre = col
        elif not col_estatus and col_upper in ["ESTATUS", "ESTADO"]:
          col_estatus = col

      for col in df.columns:
        col_upper = str(col).upper()
        if not col_ruta and "RUTA" in col_upper:
          col_ruta = col
        elif not col_rif and (
            "IDENTIFICACIÓN" in col_upper
            or "IDENTIFICACION" in col_upper
            or "RIF" in col_upper
        ):
          col_rif = col
        elif (
            not col_nombre
            and "NOMBRE" in col_upper
            and "RUTA" not in col_upper
        ):
          col_nombre = col
        elif not col_estatus and any(
            k in col_upper
            for k in ["ESTATUS", "ESTADO", "SITUACION", "CONDICION"]
        ):
          col_estatus = col

      data_dict = {}
      if col_ruta:
        data_dict["Ruta (Nombre)"] = df[col_ruta]
      if col_rif:
        data_dict["Cliente (Identificación)"] = df[col_rif]
      if col_nombre:
        data_dict["Nombre Cliente"] = df[col_nombre]
      else:
        data_dict["Nombre Cliente"] = "SIN NOMBRE"
      if col_estatus:
        data_dict["Estatus"] = df[col_estatus]

      df_limpio = pd.DataFrame(data_dict)

      if (
          "Ruta (Nombre)" not in df_limpio.columns
          or "Cliente (Identificación)" not in df_limpio.columns
      ):
        messagebox.showerror(
            "Error",
            "No se pudieron identificar las columnas requeridas (Ruta y"
            " Cliente (Identificación)) en el archivo Excel.",
        )
        return

      nombres_buscados = tuple(f"T0X{i:02d}" for i in range(1, 19))
      df_limpio["Ruta (Nombre)"] = df_limpio["Ruta (Nombre)"].astype(str)

      df_filtrado = df_limpio[
          df_limpio["Ruta (Nombre)"].str.startswith(nombres_buscados)
      ].copy()

      if df_filtrado.empty:
        messagebox.showwarning(
            "Aviso", "No se encontraron rutas coincidentes (T0X01 a T0X18)."
        )
        return

      self.ruta_archivo_actual = file_path
      self.df_filtrado_global = df_filtrado.reset_index(drop=False)

      self.df_filtrado_global["Cliente (Identificación)"] = (
          self.df_filtrado_global["Cliente (Identificación)"]
          .astype(str)
          .str.replace("-", "", regex=False)
          .str.strip()
      )

      self.frame_upload.pack_forget()

      nombre_archivo = os.path.basename(file_path)
      self.lbl_file_name.config(text=f"📄 {nombre_archivo}")

      if not self.frame_file_info.winfo_ismapped():
        self.frame_file_info.pack(
            fill="x", padx=20, pady=(4, 4), before=self.frame_checks
        )

      for widget in self.frame_checks.winfo_children():
        widget.destroy()
      self.route_vars.clear()

      rutas_counts = (
          self.df_filtrado_global["Ruta (Nombre)"].value_counts().to_dict()
      )
      rutas_encontradas = sorted(rutas_counts.keys())

      if not rutas_encontradas:
        messagebox.showwarning("Aviso", "No se encontraron rutas coincidentes.")
        self.btn_download.pack_forget()
        return

      lbl_estatus = tk.Label(
          self.frame_checks,
          text="Selecciona el/los estatus a evaluar:",
          bg="#f8f9fa",
          font=("Arial", 9, "bold"),
          fg="#333",
      )
      lbl_estatus.pack(anchor="w", padx=20, pady=2)

      frame_estatus = tk.Frame(self.frame_checks, bg="#f8f9fa")
      frame_estatus.pack(anchor="w", padx=25, pady=2)

      tk.Checkbutton(
          frame_estatus,
          text="Activos",
          variable=self.chk_activo_var,
          bg="#f8f9fa",
      ).pack(side="left", padx=5)
      tk.Checkbutton(
          frame_estatus,
          text="Inactivos",
          variable=self.chk_inactivo_var,
          bg="#f8f9fa",
      ).pack(side="left", padx=5)

      lbl_sub = tk.Label(
          self.frame_checks,
          text="Selecciona las rutas a incluir:",
          bg="#f8f9fa",
          font=("Arial", 9, "bold"),
          fg="#333",
      )
      lbl_sub.pack(anchor="w", padx=20, pady=(4, 2))

      container_scroll = tk.Frame(
          self.frame_checks, bg="#f8f9fa", bd=1, relief="solid"
      )
      container_scroll.pack(fill="both", expand=True, padx=20, pady=4)

      canvas = tk.Canvas(container_scroll, bg="white", highlightthickness=0)
      scrollbar = tk.Scrollbar(
          container_scroll, orient="vertical", command=canvas.yview
      )
      scrollable_frame = tk.Frame(canvas, bg="white")

      scrollable_frame.bind(
          "<Configure>",
          lambda e: canvas.configure(scrollregion=canvas.bbox("all")),
      )
      canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
      canvas.configure(yscrollcommand=scrollbar.set)

      canvas.pack(side="left", fill="both", expand=True)
      scrollbar.pack(side="right", fill="y")

      def _on_mousewheel(event):
        if event.delta:
          canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")
        elif event.num == 4:
          canvas.yview_scroll(-1, "units")
        elif event.num == 5:
          canvas.yview_scroll(1, "units")

      container_scroll.bind("<MouseWheel>", _on_mousewheel)
      container_scroll.bind("<Button-4>", _on_mousewheel)
      container_scroll.bind("<Button-5>", _on_mousewheel)

      canvas.bind("<MouseWheel>", _on_mousewheel)
      canvas.bind("<Button-4>", _on_mousewheel)
      canvas.bind("<Button-5>", _on_mousewheel)

      scrollable_frame.bind("<MouseWheel>", _on_mousewheel)
      scrollable_frame.bind("<Button-4>", _on_mousewheel)
      scrollable_frame.bind("<Button-5>", _on_mousewheel)

      for idx, ruta in enumerate(rutas_encontradas):
        var = tk.BooleanVar(value=True)
        var.trace_add("write", lambda *args: self.actualizar_estadisticas())
        self.route_vars[ruta] = var

        texto_ruta = f"{ruta} ({rutas_counts[ruta]})"
        chk = tk.Checkbutton(
            scrollable_frame,
            text=texto_ruta,
            variable=var,
            bg="white",
            font=("Arial", 9),
        )

        chk.bind("<MouseWheel>", _on_mousewheel)
        chk.bind("<Button-4>", _on_mousewheel)
        chk.bind("<Button-5>", _on_mousewheel)

        row = idx // 2
        col = idx % 2
        chk.grid(row=row, column=col, sticky="w", padx=10, pady=2)

      num_filas = (len(rutas_encontradas) + 1) // 2
      altura_canvas = min(130, max(40, num_filas * 22))
      canvas.config(height=altura_canvas)

      # --- PANEL DE CONTADORES / ESTADÍSTICAS INTERACTIVAS ---
      frame_stats = tk.Frame(self.frame_checks, bg="#f8f9fa")
      frame_stats.pack(fill="x", padx=20, pady=(2, 2))

      self.btn_stat_total = tk.Button(
          frame_stats,
          text="Total: 0",
          font=("Arial", 8, "bold"),
          bg="#e9ecef",
          fg="#333",
          relief="groove",
          command=lambda: self.mostrar_detalle("total"),
      )
      self.btn_stat_total.pack(side="left", padx=2)

      self.btn_stat_dup = tk.Button(
          frame_stats,
          text="Duplicados: 0",
          font=("Arial", 8, "bold"),
          bg="#fff3cd",
          fg="#856404",
          relief="groove",
          command=lambda: self.mostrar_detalle("duplicados"),
      )
      self.btn_stat_dup.pack(side="left", padx=2)

      self.btn_stat_vacio = tk.Button(
          frame_stats,
          text="Vacíos: 0",
          font=("Arial", 8, "bold"),
          bg="#f8d7da",
          fg="#721c24",
          relief="groove",
          command=lambda: self.mostrar_detalle("vacios"),
      )
      self.btn_stat_vacio.pack(side="left", padx=2)

      self.btn_stat_invalido = tk.Button(
          frame_stats,
          text="Inválidos: 0",
          font=("Arial", 8, "bold"),
          bg="#cfe2ff",
          fg="#084298",
          relief="groove",
          command=lambda: self.mostrar_detalle("invalidos"),
      )
      self.btn_stat_invalido.pack(side="left", padx=2)

      if not self.btn_download.winfo_ismapped():
        self.btn_download.pack(pady=(4, 8))

      self.actualizar_estadisticas()
      self.update_idletasks()

      messagebox.showinfo(
          "Éxito", f"Archivo cargado. {len(rutas_encontradas)} rutas listas."
      )

    except Exception as e:
      messagebox.showerror("Error", f"No se pudo leer el archivo:\n{e}")

  def obtener_df_filtrado_actual(self):
    if self.df_filtrado_global is None or self.df_filtrado_global.empty:
      return pd.DataFrame()

    rutas_seleccionadas = [
        ruta for ruta, var in self.route_vars.items() if var.get()
    ]
    df_sub = self.df_filtrado_global[
        self.df_filtrado_global["Ruta (Nombre)"].isin(rutas_seleccionadas)
    ].copy()

    activo = self.chk_activo_var.get()
    inactivo = self.chk_inactivo_var.get()

    if "Estatus" in df_sub.columns:
      s = df_sub["Estatus"].astype(str).str.strip().str.lower()
      if activo and not inactivo:
        df_sub = df_sub[
            s.str.contains("activo", case=False, na=False)
            & ~s.str.contains("inactivo", case=False, na=False)
        ]
      elif inactivo and not activo:
        df_sub = df_sub[s.str.contains("inactivo", case=False, na=False)]
      elif not activo and not inactivo:
        df_sub = df_sub.iloc[0:0]

    return df_sub

  def actualizar_estadisticas(self):
    df_sub = self.obtener_df_filtrado_actual()
    if df_sub.empty:
      if hasattr(self, "btn_stat_total"):
        self.btn_stat_total.config(text="Total: 0")
        self.btn_stat_dup.config(text="Duplicados: 0")
        self.btn_stat_vacio.config(text="Vacíos: 0")
        self.btn_stat_invalido.config(text="Inválidos: 0")
      return

    rifs_serie = df_sub["Cliente (Identificación)"].astype(str).str.strip()
    mask_vacios = (
        df_sub["Cliente (Identificación)"].isna()
        | (rifs_serie == "")
        | (rifs_serie.str.lower().isin(["nan", "none", "null", "nat", "0"]))
    )
    vacios_count = mask_vacios.sum()

    df_validos = df_sub[~mask_vacios]
    dups_mask = df_validos.duplicated(
        subset=["Cliente (Identificación)"], keep=False
    )
    dups_count = dups_mask.sum()

    def es_rif_valido(rif):
      rif_limpio = re.sub(r"[^0-9A-Za-z]", "", str(rif))
      return bool(re.match(r"^[VEJPGvejpg]\d{7,10}$", rif_limpio))

    invalidos_count = sum(
        1
        for rif in df_validos["Cliente (Identificación)"]
        if not es_rif_valido(rif)
    )

    total_unicos = df_validos["Cliente (Identificación)"].nunique()

    if hasattr(self, "btn_stat_total"):
      self.btn_stat_total.config(text=f"Total: {total_unicos}")
      self.btn_stat_dup.config(text=f"Duplicados: {dups_count}")
      self.btn_stat_vacio.config(text=f"Vacíos: {vacios_count}")
      self.btn_stat_invalido.config(text=f"Inválidos: {invalidos_count}")

  def mostrar_detalle(self, tipo):
    df_sub = self.obtener_df_filtrado_actual()
    if df_sub.empty:
      messagebox.showinfo("Aviso", "No hay datos para mostrar con el filtro actual.")
      return

    rifs_serie = df_sub["Cliente (Identificación)"].astype(str).str.strip()
    mask_vacios = (
        df_sub["Cliente (Identificación)"].isna()
        | (rifs_serie == "")
        | (rifs_serie.str.lower().isin(["nan", "none", "null", "nat", "0"]))
    )

    if tipo == "vacios":
      df_resultado = df_sub[mask_vacios]
      titulo = "Registros con celdas de RIF Vacías o Nulas"
    elif tipo == "duplicados":
      df_validos = df_sub[~mask_vacios]
      dups_mask = df_validos.duplicated(
          subset=["Cliente (Identificación)"], keep=False
      )
      df_resultado = df_validos[dups_mask].sort_values(
          by="Cliente (Identificación)"
      )
      titulo = "Registros con RIFs Duplicados (Agrupados)"
    elif tipo == "invalidos":

      def es_rif_valido(rif):
        rif_limpio = re.sub(r"[^0-9A-Za-z]", "", str(rif))
        return bool(re.match(r"^[VEJPGvejpg]\d{7,10}$", rif_limpio))

      df_validos = df_sub[~mask_vacios]
      invalidos_mask = [
          not es_rif_valido(rif)
          for rif in df_validos["Cliente (Identificación)"]
      ]
      df_resultado = df_validos[invalidos_mask]
      titulo = "Registros con RIFs de Formato Inválido"
    else:
      df_validos = df_sub[~mask_vacios]
      df_resultado = df_validos.drop_duplicates(
          subset=["Cliente (Identificación)"]
      ).sort_values(by="Cliente (Identificación)")
      titulo = "Total de Clientes Únicos a Consultar"

    if df_resultado.empty:
      messagebox.showinfo("Información", f"No se encontraron elementos en la categoría '{tipo}'.")
      return

    top = tk.Toplevel(self)
    top.title(titulo)
    top.geometry("700x400")
    top.config(bg="#f8f9fa")

    lbl_titulo = tk.Label(
        top,
        text=f"📋 {titulo} ({len(df_resultado)} encontrados)",
        font=("Arial", 11, "bold"),
        bg="#f8f9fa",
        fg="#333",
    )
    lbl_titulo.pack(anchor="w", padx=15, pady=10)

    frame_tree = tk.Frame(top, bg="#f8f9fa")
    frame_tree.pack(fill="both", expand=True, padx=15, pady=(0, 15))

    tree = ttk.Treeview(
        frame_tree,
        columns=("fila", "cliente", "rif", "ruta", "estatus"),
        show="headings",
        height=12,
    )
    tree.heading("fila", text="Fila Excel")
    tree.heading("cliente", text="Nombre del Cliente")
    tree.heading("rif", text="RIF / Identificación")
    tree.heading("ruta", text="Ruta")
    tree.heading("estatus", text="Estatus")

    tree.column("fila", width=70, anchor="center")
    tree.column("cliente", width=260, anchor="w")
    tree.column("rif", width=120, anchor="center")
    tree.column("ruta", width=90, anchor="center")
    tree.column("estatus", width=80, anchor="center")

    scrollbar_tree = ttk.Scrollbar(
        frame_tree, orient="vertical", command=tree.yview
    )
    tree.configure(yscrollcommand=scrollbar_tree.set)

    tree.pack(side="left", fill="both", expand=True)
    scrollbar_tree.pack(side="right", fill="y")

    for _, row in df_resultado.iterrows():
      idx_orig = row.get("index", 0)
      fila_excel = self.excel_header_row + 2 + idx_orig
      nombre_cli = row.get("Nombre Cliente", "SIN NOMBRE")
      rif_val = row.get("Cliente (Identificación)", "")
      ruta_val = row.get("Ruta (Nombre)", "")
      estatus_val = row.get("Estatus", "")

      tree.insert(
          "",
          "end",
          values=(fila_excel, nombre_cli, rif_val, ruta_val, estatus_val),
      )

  def escribir_log(self, mensaje):
    self.after(0, self._insertar_texto, mensaje)

  def _insertar_texto(self, mensaje):
    self.txt_log.config(state="normal")
    self.txt_log.insert(tk.END, mensaje + "\n")
    self.txt_log.see(tk.END)
    self.txt_log.config(state="disabled")

  def solicitar_cancelacion(self):
    """Muestra un diálogo de confirmación para cancelar el proceso."""
    respuesta = messagebox.askyesno(
        "Confirmar Cancelación",
        "¿Seguro que quiere cancelar la consulta?",
        parent=self.top_log if hasattr(self, 'top_log') and self.top_log.winfo_exists() else self
    )
    if respuesta:
      self.cancelar_proceso = True
      if hasattr(self, 'top_log') and self.top_log.winfo_exists():
        self.top_log.destroy()
      self.restaurar_boton_iniciar()

  def restaurar_boton_iniciar(self):
    """Devuelve el botón a su estado original de 'Iniciar Consulta'."""
    self.btn_download.config(
        text="Iniciar Consulta",
        bg="#28a745",
        command=self.iniciar_hilo_consultas,
        state="normal"
    )

  def iniciar_hilo_consultas(self):
    df_final = self.obtener_df_filtrado_actual()
    if df_final.empty:
      messagebox.showwarning(
          "Atención", "No hay registros seleccionados para procesar."
      )
      return

    rifs_a_consultar = (
        df_final["Cliente (Identificación)"].dropna().unique().tolist()
    )

    if not rifs_a_consultar:
      messagebox.showwarning("Aviso", "No hay RIFs válidos para consultar.")
      return

    # Solicitar ruta de guardado antes de iniciar el proceso
    base, ext = os.path.splitext(self.ruta_archivo_actual)
    nombre_sugerido = f"{os.path.basename(base)}_PROCESADO{ext}"
    self.ruta_guardado = filedialog.asksaveasfilename(
        title="Guardar archivo procesado como...",
        initialfile=nombre_sugerido,
        defaultextension=".xlsx",
        filetypes=[("Archivos de Excel", "*.xlsx"), ("Todos los archivos", "*.*")]
    )

    # Si el usuario cierra el cuadro de diálogo sin elegir ruta, abortar
    if not self.ruta_guardado:
      return

    self.cancelar_proceso = False

    self.top_log = tk.Toplevel(self)
    self.top_log.title("Consola SENIAT - Procesamiento por Lotes")
    self.top_log.geometry("650x450")
    self.top_log.config(bg="#1e1e1e")

    # Vincular cierre de la ventana de la consola (la 'X') con la cancelación
    self.top_log.protocol("WM_DELETE_WINDOW", self.solicitar_cancelacion)

    self.txt_log = tk.Text(
        self.top_log, bg="#1e1e1e", fg="#4af626", font=("Consolas", 10)
    )
    self.txt_log.pack(fill="both", expand=True, padx=10, pady=10)
    self.txt_log.config(state="disabled")

    # Modificar el botón para que actúe como "Cancelar"
    self.btn_download.config(
        text="Cancelar Consulta",
        bg="#dc3545",
        command=self.solicitar_cancelacion,
        state="normal"
    )

    hilo_maestro = threading.Thread(
        target=self.gestor_multihilo_background,
        args=(rifs_a_consultar,),
        daemon=True,
    )
    hilo_maestro.start()

  def procesar_lote_hilo(
      self,
      id_hilo,
      chunk_rifs,
      mapa_contribuyentes,
      lock,
      contador_progreso,
      total_rifs,
      api
  ):
    self.escribir_log(
        f"[Hilo #{id_hilo}] Iniciado. Lote asignado: {len(chunk_rifs)} RIFs."
    )

    for rif in chunk_rifs:
      # --- Interrumpir ciclo si el proceso fue cancelado ---
      if getattr(self, "cancelar_proceso", False):
        break

      max_intentos = 12
      intento = 1
      res = None

      while intento <= max_intentos:
        res = api.consultar_rif(rif)

        if res and res.get("status") in ["success", "not_found"]:
          if res.get("status") == "success":
            self.escribir_log(f"[{rif}] ✅ Captcha Correcto.")
          break

        if res and res.get("status") == "fail_ocr":
          self.escribir_log(
              f"[{rif}] ⚠️ Captcha Incorrecto (Intento"
              f" {intento}/{max_intentos}). Reintentando..."
          )
          intento += 1
          time.sleep(1.0)
          continue

        self.escribir_log(f"[{rif}] ❌ Error de conexión o servidor.")
        break

      # --- Lógica de asignación de estatus según resultado o agotamiento de intentos ---
      if res and res.get("status") == "success":
        nombre = res.get("nombre", "SIN NOMBRE").title()
        condicion = res.get("condicion", "").upper()

        if "AGENTE DE RETENCIÓN DEL IVA" in condicion:
          self.escribir_log(f"[{rif}] ▶ Procesado: {nombre} | Especial: SI")
          es_contribuyente = "SI"
        else:
          self.escribir_log(f"[{rif}] ▶ Procesado: {nombre} | Especial: NO")
          es_contribuyente = "NO"
      elif res and res.get("status") == "not_found":
        self.escribir_log(f"[{rif}] ▶ El RIF No Existe en SENIAT.")
        es_contribuyente = "NO EXISTE"
      else:
        self.escribir_log(f"[{rif}] ⚠️ Excedido número de reintentos ({max_intentos}). Marcar como REVISAR.")
        es_contribuyente = "REVISAR"

      with lock:
        mapa_contribuyentes[rif] = es_contribuyente
        contador_progreso[0] += 1
        completados = contador_progreso[0]

      if completados % 10 == 0 or completados == total_rifs:
        self.escribir_log(
            f"\n--- PROGRESO GENERAL: {completados} / {total_rifs} COMPLETADOS"
            " ---\n"
        )

      time.sleep(0.5)

  def gestor_multihilo_background(self, rifs_a_consultar):
    total = len(rifs_a_consultar)
    MAX_WORKERS = 8

    self.escribir_log(
        f"--- INICIANDO DISTRIBUCIÓN POR LOTES ({total} REGISTROS) ---"
    )
    self.escribir_log(f"--- HILOS ACTIVOS: {MAX_WORKERS} ---\n")
    self.escribir_log(
        "--- Inicializando motor de reconocimiento visual... ---\n"
    )

    try:
        motor_ocr_global = SeniatOCR()
    except Exception as e:
        self.escribir_log(f"ERROR CRITICO: Fallo al inicializar OCR: {e}")
        return

    k, m = divmod(total, MAX_WORKERS)
    chunks = [
        rifs_a_consultar[i * k + min(i, m) : (i + 1) * k + min(i + 1, m)]
        for i in range(MAX_WORKERS)
    ]
    chunks = [c for c in chunks if len(c) > 0]

    mapa_contribuyentes = {}
    lock = threading.Lock()
    contador_progreso = [0]

    with ThreadPoolExecutor(max_workers=len(chunks)) as executor:
      futuros = []
      for idx, chunk in enumerate(chunks, 1):
        futuro = executor.submit(
            self.procesar_lote_hilo,
            idx,
            chunk,
            mapa_contribuyentes,
            lock,
            contador_progreso,
            total,
            motor_ocr_global
        )
        futuros.append(futuro)

      wait(futuros)

      for f in futuros:
          try:
              f.result()
          except Exception as e:
              self.escribir_log(f"ERROR FATAL EN HILO: {e}")

    # --- Evitar generar el documento si se canceló la operación ---
    if getattr(self, "cancelar_proceso", False):
      return

    self.generar_excel_final(mapa_contribuyentes)

  def generar_excel_final(self, mapa_contribuyentes):
    self.escribir_log("\n--- CONSULTAS FINALIZADAS ---")
    self.escribir_log("Inyectando resultados en el documento Excel...")
    try:
      wb = openpyxl.load_workbook(self.ruta_archivo_actual)
      ws = wb.active

      header_row = None
      col_rif = None
      col_contrib_existente = None

      for r in range(1, min(30, ws.max_row + 1)):
        for c in range(1, ws.max_column + 1):
          val = str(ws.cell(row=r, column=c).value or "").strip()
          if val == "Cliente (Identificación)":
            header_row = r
            col_rif = c
          elif val in ["Contribuyente", "Contribuyente Especial"]:
            col_contrib_existente = c

      if col_rif is not None:
        if col_contrib_existente is not None:
          col_dest = col_contrib_existente
        else:
          col_dest = ws.max_column + 1
          letra_nueva_col = get_column_letter(col_dest)

          celda_header = ws.cell(
              row=header_row, column=col_dest, value="Contribuyente"
          )
          celda_referencia = ws.cell(row=header_row, column=col_dest - 1)

          if celda_referencia.has_style:
            celda_header.font = copy(celda_referencia.font)
            celda_header.border = copy(celda_referencia.border)
            celda_header.fill = copy(celda_referencia.fill)
            celda_header.alignment = copy(celda_referencia.alignment)

          ws.column_dimensions[letra_nueva_col].width = 16

        for row in range(header_row + 1, ws.max_row + 1):
          valor_rif = (
              str(ws.cell(row=row, column=col_rif).value or "")
              .replace("-", "")
              .strip()
          )
          if valor_rif and valor_rif in mapa_contribuyentes:
            celda_resultado = ws.cell(
                row=row, column=col_dest, value=mapa_contribuyentes[valor_rif]
            )
            celda_resultado.alignment = openpyxl.styles.Alignment(
                horizontal="center"
            )

        # --- EXPANDIR O CREAR AUTOFILTRO DE EXCEL ---
        col_final_letra = get_column_letter(ws.max_column)
        if ws.auto_filter and ws.auto_filter.ref:
          ref_actual = str(ws.auto_filter.ref)
          if ":" in ref_actual:
            inicio, fin = ref_actual.split(":")
            filas_fin = re.findall(r"\d+", fin)
            num_fila_fin = filas_fin[0] if filas_fin else str(ws.max_row)
            ws.auto_filter.ref = f"{inicio}:{col_final_letra}{num_fila_fin}"
          else:
            ws.auto_filter.ref = f"A{header_row}:{col_final_letra}{ws.max_row}"
        else:
          ws.auto_filter.ref = f"A{header_row}:{col_final_letra}{ws.max_row}"

        nuevo_archivo = self.ruta_guardado
        wb.save(nuevo_archivo)

        self.escribir_log(
            f"¡ÉXITO! Archivo generado correctamente:\n{nuevo_archivo}"
        )
        self.after(
            0,
            lambda: messagebox.showinfo(
                "Completado", f"Archivo procesado guardado:\n\n{nuevo_archivo}"
            ),
        )
      else:
        self.escribir_log("Error: No se encontró la columna de identificación.")

    except Exception as e:
      self.escribir_log(f"Error fatal guardando Excel: {e}")
    finally:
      self.after(0, self.restaurar_boton_iniciar)


if __name__ == "__main__":
  app = ExcelUploaderApp()
  app.mainloop()